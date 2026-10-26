from dotenv import load_dotenv
from fastapi import FastAPI, UploadFile, File, Form
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
import uvicorn
from io import BytesIO
from ocr_utils import extract_text_from_file
from preeti import normalize_nepali_text, preeti_to_unicode, detect_preeti
from translator import translate_text
from docx import Document
from fpdf import FPDF

load_dotenv()

app = FastAPI(title="Nepali OCR + Translate API")

# serve static UI under /ui so API routes are not intercepted
app.mount('/ui', StaticFiles(directory='static', html=True), name='static')

from fastapi.responses import RedirectResponse


@app.get('/')
def root():
    return RedirectResponse(url='/ui/')


@app.post('/extract')
async def extract(file: UploadFile = File(...), preeti: str = Form('false')):
    try:
        data = await file.read()
        text = extract_text_from_file(data, filename=file.filename)
        # Always normalize mixed Nepali inputs before returning or translating.
        if preeti.lower() == 'true':
            text = preeti_to_unicode(text)
        elif preeti.lower() == 'auto':
            if detect_preeti(text):
                text = preeti_to_unicode(text)
        text = normalize_nepali_text(text)
        return JSONResponse({"filename": file.filename, "text": text})
    except ValueError as exc:
        return JSONResponse({"detail": str(exc)}, status_code=400)
    except RuntimeError as exc:
        return JSONResponse({"detail": str(exc)}, status_code=500)
    except Exception as exc:
        return JSONResponse({"detail": f'Extraction failed: {exc}'}, status_code=500)


@app.post('/translate')
async def translate(text: str = Form(...), src: str = Form('ne'), dest: str = Form('en')):
    normalized_text = normalize_nepali_text(text)
    translated = translate_text(normalized_text, src=src, dest=dest)
    status_code = 200
    if translated.startswith('[translation-unavailable]'):
        status_code = 503
    return JSONResponse({"translated_text": translated}, status_code=status_code)


def docx_bytes_from_text(title: str, original: str, translated: str):
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_heading('Original (Nepali)', level=2)
    doc.add_paragraph(original)
    doc.add_heading('Translated (English)', level=2)
    doc.add_paragraph(translated)
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


def sanitize_for_pdf(text: str) -> str:
    return text.encode('latin-1', 'replace').decode('latin-1')


def pdf_bytes_from_text(title: str, original: str, translated: str):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font('Arial', 'B', 16)
    pdf.multi_cell(0, 8, sanitize_for_pdf(title))
    pdf.ln(4)
    pdf.set_font('Arial', '', 12)
    pdf.multi_cell(0, 7, 'Original (Nepali)\n')
    pdf.set_font('Arial', '', 11)
    pdf.multi_cell(0, 6, sanitize_for_pdf(original))
    pdf.ln(4)
    pdf.set_font('Arial', '', 12)
    pdf.multi_cell(0, 7, 'Translated (English)\n')
    pdf.set_font('Arial', '', 11)
    pdf.multi_cell(0, 6, sanitize_for_pdf(translated))
    bio = BytesIO()
    payload = pdf.output(dest='S')
    if isinstance(payload, str):
        payload = payload.encode('latin-1')
    bio.write(payload)
    bio.seek(0)
    return bio


@app.post('/download')
async def download(
    title: str = Form('Translation'),
    original: str = Form(...),
    translated: str = Form(...),
    format: str = Form('pdf'),
):
    if format.lower() == 'docx':
        bio = docx_bytes_from_text(title, original, translated)
        return StreamingResponse(bio, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document', headers={
            'Content-Disposition': f'attachment; filename="{title}.docx"'
        })
    else:
        bio = pdf_bytes_from_text(title, original, translated)
        return StreamingResponse(bio, media_type='application/pdf', headers={
            'Content-Disposition': f'attachment; filename="{title}.pdf"'
        })


if __name__ == '__main__':
    uvicorn.run(app, host='0.0.0.0', port=8000)
